from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import os
from typing import Any, Iterable

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from business_cycle import (
    BUSINESS_CYCLE_MODEL_VERSION,
    ECONOMY_REGIME_MODEL_VERSION,
    INFLATION_LAYER_MODEL_VERSION,
)
from funding_conditions import MODEL_VERSION as FUNDING_MODEL_VERSION
from global_dashboard import classify_cross_cycle
from liquidity_forecast import BREADTH_VERSION, FORECAST_VERSION
from macro_research_export import SeriesMeta, write_dataframe_sheet
from macro_surprises import MACRO_SURPRISES_MODEL_VERSION
from market_cycle import MARKET_CYCLE_MODEL_VERSION
from rates_financial_conditions import MODEL_VERSION as RATES_FC_MODEL_VERSION
from treasury_fiscal_regime import MODEL_VERSION as TREASURY_FISCAL_MODEL_VERSION


DEFAULT_START_DATE = "2010-01-01"
EXPORT_FILENAME = "financial_fragility_validation_export.xlsx"

HORIZONS = {"1M": 4, "3M": 13, "6M": 26, "12M": 52}

GLOBAL_LIQUIDITY_FIELDS: dict[str, tuple[str, ...]] = {
    "GlobalLiquidityState": ("final_regime_label", "impulse_state"),
    "GlobalLiquidityDirection": ("direction_13w_state",),
    "LiquidityCyclePhase": ("long_cycle_phase",),
    "GlobalLiquidityLevel": ("global_liquidity_score",),
    "CBImpulse": ("cb_impulse",),
    "USNetLiquidityImpulse": ("usnl_impulse",),
    "GlobalM2_13W": ("m2_13w",),
    "GlobalM2_26W": ("m2_26w",),
    "GlobalM2_52W": ("m2_52w",),
}

TRANSITION_FIELDS: dict[str, tuple[str, ...]] = {
    "FastTransitionRisk": ("Fast_Transition_Risk",),
    "MacroTransitionRisk": ("Macro_Transition_Risk",),
    "PositioningVulnerability": ("PositioningRisk",),
    "PositioningVulnerabilityState": ("PositioningState",),
}

MARKET_FIELDS: dict[str, tuple[str, ...]] = {
    "SPX_Close": ("SPX_Close",),
    "SPY_Close": ("CurrentRiskSPYRawClose", "SPY", "SPY_Close"),
    "StructuralMarketCyclePhase": ("StructuralMarketCyclePhase",),
    "StructuralSMA_ROC_Phase": ("Structural_SMA_ROC_Phase",),
    "StructuralExtensionPct": ("StructuralExtensionPct",),
    "StructuralExtensionPercentile_PIT": ("StructuralExtensionPercentile",),
    "SPX_ROC36M": ("SPX_ROC36M",),
    "SPX_ROC36M_3MMA": ("SPX_ROC36M_3MMA",),
    "StructuralROCMomentum": ("StructuralROCMomentum",),
    "StructuralROCMomentumState": ("StructuralROCMomentumState",),
    "StructuralExtensionDirection12M": ("StructuralExtensionDirection_12M",),
    "StructuralCycleAgeMonths": ("StructuralCycleAgeMonths",),
    "StructuralCycleAgePercentile": ("StructuralAgePctOfHistoricalMedian",),
    "StructuralMaturity": ("StructuralMaturityStatus", "StructuralMaturity"),
    "SPXMomentumCyclePhase": ("MomentumCyclePhase",),
    "SPX_ROC12M": ("SPX_ROC12M",),
    "SPX_ROC12M_3MMA": ("SPX_ROC12M_3MMA",),
    "ROCMomentum3M": ("SPX_ROC_Momentum_3M",),
    "SMA200WExtensionPct": ("SMA200WExtensionPct",),
    "SMA200WExtensionPercentile_PIT": ("SMA200WExtensionPercentile",),
    "SMA200WExtensionPhase": ("SMA200WZone",),
    "MediumTerm_SMA_ROC_Phase": ("MediumTerm_SMA_ROC_Phase",),
    "PrimaryMarketCycleValue": ("PrimaryMarketCycleValue",),
    "LongMarketExtensionCycleValue": ("LongExtensionCycleValue", "LongMarketExtensionCycleValue"),
    "MonthsSincePrimaryCycleTrough": ("PrimaryCycleMonthsSinceTrough", "MonthsSincePrimaryCycleTrough"),
    "PositioningVulnerabilityState": ("PositioningVulnerability",),
    "AAII_Bearish_Percentile": ("AAII_Bearish_Percentile",),
    "VIX_AssetManager_Percentile": ("VIX_AssetManager_Percentile",),
    "CurrentMarketRiskState": ("CurrentMarketRiskState", "HistoricalCurrentRiskState"),
    "CurrentMarketRiskDirection": ("CurrentMarketRiskDirection",),
    "CurrentMarketRiskScore": ("HistoricalCurrentRiskScore", "CurrentRiskComponentAverage"),
    "DrawdownRisk": ("HistoricalDrawdownRisk", "CurrentRiskDrawdownRisk"),
    "HighBetaRisk": ("HistoricalHighBetaRisk", "CurrentRiskHighBetaRisk"),
    "BreadthRisk": ("HistoricalBreadthRisk", "CurrentRiskBreadthRisk", "BreadthRisk"),
    "RSIDivergenceRisk": ("HistoricalRSIDivergenceRisk", "CurrentRiskRSIDivergenceRisk"),
    "CombinedVIXRisk": ("HistoricalCombinedVIXRisk", "CurrentRiskVIXRisk"),
    "VIXLevelRisk": ("HistoricalVIXLevelRisk", "CurrentRiskVIXRisk"),
    "VIXMomentumRisk": ("HistoricalVIXMomentumRisk", "CurrentRiskVIXRisk"),
    "VIXTermStructureRisk": ("HistoricalVIXTermStructureRisk", "CurrentRiskVIXRisk"),
    "RealizedVolRisk": ("RealizedVolRisk", "RealizedVolatilityRisk"),
    "VIX": ("VIX",),
    "VIX3M": ("VIX3M",),
    "VIX_VIX3M_Ratio": ("VIXTermStructure", "VIX_VIX3M_Ratio"),
    "VIX_20D_Change": ("VIX20DChange", "VIX_20D_Change"),
    "RealizedVolatility_20D": ("RealizedVol20D", "SPX_RealizedVol20D", "RealizedVolatility_20D"),
    "QQQ_SPX_RS13W": ("QQQ_SPX_RS_13W",),
    "QQQ_SPX_RS26W": ("QQQ_SPX_RS_26W",),
    "QQQ_SPX_RS13W_Percentile": ("QQQ_SPX_RS_13W_Pctl",),
    "QQQ_SPX_RS26W_Percentile": ("QQQ_SPX_RS_26W_Pctl",),
    "BTC_SPX_RS13W": ("BTC_SPX_RS_13W",),
    "BTC_SPX_RS26W": ("BTC_SPX_RS_26W",),
    "BTC_SPX_RS13W_Percentile": ("BTC_SPX_RS_13W_Pctl",),
    "BTC_SPX_RS26W_Percentile": ("BTC_SPX_RS_26W_Pctl",),
    "BreadthAbove50D": ("SPXAboveSMA50D",),
    "BreadthAbove200D": ("SPXAboveSMA200D",),
    "BreadthLevelRisk": ("BreadthLevelRisk",),
    "BreadthMomentumRisk": ("BreadthMomentumRisk",),
    "RSIDivergenceActive": ("RSIDivergenceActive",),
    "RSIDivergenceDurationWeeks": ("RSIDivergenceDurationWeeks",),
    "RSIDropPoints": ("RSIDropPoints",),
    "PriceGainDuringDivergence": ("PriceGainDuringDivergence",),
    "MarketCycle_LastUpdated": ("LastUpdated",),
    "CurrentRisk_LastUpdated": ("LastUpdated",),
}

TREASURY_FIELDS: dict[str, tuple[str, ...]] = {
    "TreasuryLiquidityState": ("TreasuryLiquidityState",),
    "TreasuryLiquidityImpulse": ("TreasuryLiquidityImpulse",),
    "TreasuryLiquidityImpulse_Fast": ("FastLiquidityImpulse",),
    "TreasuryLiquidityImpulse_Medium": ("MediumLiquidityImpulse",),
    "TreasuryLiquidityImpulse_Slow": ("SlowLiquidityImpulse",),
    "FedImpulse": ("FedImpulse_13W",),
    "TGAImpulse": ("TGAImpulse_13W",),
    "RRPImpulse": ("RRPImpulse_13W",),
    "WTREGEN": ("WTREGEN",),
    "WDTGAL": ("WDTGAL",),
    "RRPONTSYD": ("RRP", "RRPONTSYD"),
    "WALCL": ("WALCL",),
    "WRESBAL": ("WRESBAL",),
    "FiscalGrowthImpulse": ("FiscalImpulse",),
    "FiscalGrowthState": ("FiscalImpulseState",),
    "FiscalStanceRaw": ("FiscalStanceRaw",),
    "FiscalStanceState": ("FiscalStanceState",),
    "FastFiscalImpulseRaw": ("FastFiscalImpulseRaw",),
    "MediumFiscalImpulseRaw": ("MediumFiscalImpulseRaw",),
    "StructuralFiscalImpulseRaw": ("StructuralFiscalImpulseRaw",),
    "FastFiscalZ": ("FastFiscalZ",),
    "MediumFiscalZ": ("MediumFiscalZ",),
    "StructuralFiscalZ": ("StructuralFiscalZ",),
    "SpendingImpulse": ("SpendingImpulse",),
    "RevenueImpulse": ("RevenueImpulse",),
    "DeficitGDPPercentile": ("FiscalStancePercentile",),
    "TreasuryFinancingPressure": ("TreasuryFinancingPressure",),
    "FinancingPressureScore": ("AbsorptionTightness",),
    "TotalSupplyPercentile": ("TotalSupplyPercentile",),
    "BillSupplyPercentile": ("BillSupplyPercentile",),
    "DurationSupplyPercentile": ("DurationSupplyPercentile",),
    "BillFinancingShare": ("BillFinancingShare",),
    "DurationSupplyProxy": ("DurationSupplyPercentile",),
    "TotalMarketableTreasury": ("FGTSL",),
    "BillSupply": ("Bills",),
    "NonBillSupply": ("NonBillDebt",),
    "AbsorptionState": ("AbsorptionCapacity",),
    "PolicyMix": ("PolicyMix",),
}

MACRO_SURPRISE_FIELDS: dict[str, tuple[str, ...]] = {
    "GrowthSurpriseScore": ("GrowthSurpriseScore",),
    "GrowthSurpriseState": ("GrowthSurpriseState",),
    "InflationSurpriseScore": ("InflationSurpriseScore",),
    "InflationSurpriseState": ("InflationSurpriseState",),
    "TurningSignal": ("TurningSignal",),
    "ISM_Surprise_Z": ("PMI_Surprise_Z", "ISM_Surprise_Z"),
    "RetailSales_Surprise_Z": ("RetailSales_Surprise_Z",),
    "InitialClaims_Surprise_Z": ("InitialJoblessClaims_Surprise_Z", "InitialClaims_Surprise_Z"),
    "CPI_Surprise_Z": ("CPI_Surprise_Z",),
}

BASE_ALIASES: dict[str, tuple[str, ...]] = {
    "GlobalLiquidityState": ("GlobalLiquidity_State",),
    "GlobalLiquidityDirection": ("GlobalLiquidity_Direction_State",),
    "LiquidityCyclePhase": ("LongLiquidityCycle_Phase",),
    "LiquidityPressure": ("LiquidityPressureScore",),
    "PolicyResponse": ("PolicyResponseScore",),
    "LiquidityBreadthState": ("BreadthParticipationState",),
    "RiskReductionState": ("RiskReductionWarning",),
    "GlobalLiquidityLevel": ("GlobalLiquidityScore",),
    "GlobalLiquidityROC": ("GlobalLiquidityScore_Change_4W",),
    "GlobalLiquidityROC_13W": ("GlobalLiquidityScore_Change_13W",),
    "GlobalLiquidityROC_26W": ("GlobalLiquidityScore_Change_26W",),
    "LiquidityPressureScore_Raw": ("LiquidityPressureScore",),
    "PolicyResponseScore_Raw": ("PolicyResponseScore",),
    "LiquidityForecastSignal": ("LiquidityForwardSignal",),
    "USNetLiquidityImpulse": ("USNLImpulse",),
    "RatesPressure": ("RatesPressureScore",),
    "US2Y": ("DGS2",),
    "US10YRealYield": ("DFII10",),
    "US10YRealYieldMomentum": ("RealYieldMomentum",),
    "FedFundsRate": ("FEDFUNDS",),
    "CoreFCLevel": ("FinancialConditionsLevel",),
    "CoreFCDirection": ("FinancialConditionsDirection",),
    "CoreFCState": ("FinancialConditionsStressLevel",),
    "FCStressOverlay": ("FCConfirmationStatus",),
    "NFCI": ("NFCILevel",),
    "ANFCI": ("ANFCILevel",),
    "DXY": ("DXY_Close", "DXY"),
    "MOVE": ("MOVE_Close", "MOVE"),
    "MOVE_Z": ("MOVE_Level",),
    "HYOAS": ("HY_OAS",),
    "IGOAS": ("IG_OAS",),
    "IG_OAS_Z": ("IG_OAS_Z",),
    "PositioningVulnerability": ("PositioningRisk",),
    "BusinessCyclePhase": ("BusinessCycleState",),
    "InflationDirection": ("InflationState",),
    "InflationLeadingScore": ("InflationDirectionScore",),
    "MarketPricingInflationScore": ("MarketPricingScore",),
    "ModelInflationScore": ("ModelImpliedInflationScore",),
    "InflationTransitionFlag": ("InflationTransitionZone",),
    "Michigan1Y": ("MICH",),
    "CorePCE": ("PCEPILFE",),
    "CoreCPI": ("CPILFESL",),
    "HeadlineCPI": ("CPIAUCSL",),
    "PPI": ("PPIACO",),
}

LIQUIDITY_FORECAST_FIELDS: dict[str, tuple[str, ...]] = {
    "LiquidityForecastState": ("LiquidityForecastState",),
    "LiquidityPressure": ("LiquidityPressureScore",),
    "PolicyResponse": ("PolicyResponseScore",),
    "LiquidityPressureScore_Raw": ("LiquidityPressureScore",),
    "PolicyResponseScore_Raw": ("PolicyResponseScore",),
    "LiquidityForecastSignal": ("LiquidityForwardSignal",),
    "RiskOnState": ("RiskOnState",),
    "RiskReductionState": ("RiskReductionWarning",),
    "LiquidityBreadthState": ("BreadthParticipationState",),
    "CBImpulse": ("CBImpulse",),
    "USNetLiquidityImpulse": ("USNLImpulse",),
    "BankReservesImpulse": ("BankReservesImpulse",),
    "DXY_13W_Change": ("DXY_13W_Change",),
    "MOVE_Level_LiquidityPressure": ("MOVE",),
    "US2Y_LiquidityPressure": ("US2Y",),
    "ISM_LiquidityPressure": ("ISM_Manufacturing",),
    "CFNAI_LiquidityPressure": ("CFNAI",),
    "ContinuingClaims_13W_Change": ("ContinuingClaims_13W_Change",),
    "US10YTermPremium_26W_Change": ("US10Y_TermPremium_26W_Change",),
}

RATES_FIELDS: dict[str, tuple[str, ...]] = {
    "RatesPressure": ("RatesPressureScore",),
    "RatesPressureScore": ("RatesPressureScore",),
    "US2Y": ("DGS2",),
    "US10YRealYield": ("DFII10",),
    "US2YMomentum": ("US2YMomentum",),
    "US10YRealYieldMomentum": ("RealYieldMomentum",),
    "FedFundsRate": ("FEDFUNDS",),
    "CoreFCLevel": ("FinancialConditionsLevel",),
    "CoreFCDirection": ("FinancialConditionsDirection",),
    "CoreFCState": ("FinancialConditionsStressLevel",),
    "FCStressOverlay": ("FCConfirmationStatus",),
    "NFCI": ("NFCI", "NFCILevel"),
    "ANFCI": ("ANFCI", "ANFCILevel"),
    "DXY": ("DXY",),
    "MOVE": ("MOVE",),
    "MOVE_Z": ("MOVE_Level",),
    "HYOAS": ("BAMLH0A0HYM2",),
    "IGOAS": ("BAMLC0A0CM",),
    "HY_OAS_Z": ("HY_OAS_Z",),
    "IG_OAS_Z": ("IG_OAS_Z",),
    "HY_OAS_Momentum": ("HY_OAS_Momentum",),
    "IG_OAS_Momentum": ("IG_OAS_Momentum",),
    "CreditLevel": ("CreditLevel",),
    "CreditDirection": ("CreditDirection",),
}

FUNDING_FIELDS: dict[str, tuple[str, ...]] = {
    "FundingState": ("FundingState",),
    "FundingDirection": ("FundingDirection",),
    "FundingCore": ("FundingCore",),
    "MoneyMarketStress": ("MoneyMarketStress",),
    "ReservePressure": ("ReservePressure",),
    "ReserveVulnerability": ("ReserveVulnerability",),
    "ReserveDrain": ("ReserveDrain",),
    "PersistentFundingFlag": ("PersistentFundingFlag",),
    "TechnicalFundingFlag": ("TechnicalFundingFlag",),
    "CalendarTechnicalFlag": ("CalendarTechnicalFlag",),
    "ReserveVulnerabilityWatch": ("ReserveVulnerabilityWatch",),
    "CollateralStress": ("CollateralStress",),
    "SOFR99": ("SOFR99",),
    "DFF": ("DFF",),
    "IORB": ("IORB",),
    "FundingSpreadRaw": ("FundingSpreadRaw",),
    "FundingSpreadSmooth": ("FundingSpreadSmooth",),
    "FundingSpreadRobustZ": ("FundingSpreadRobustZ",),
    "WRESBAL_Funding": ("WRESBAL",),
    "ReserveGDP": ("ReserveGDP",),
    "ReserveGDP_Z": ("ReserveGDP_Z",),
    "ReserveChange13W": ("ReserveChange13W",),
    "ReserveChange13W_Z": ("ReserveChange13W_Z",),
    "MOVE_Funding": ("MOVE",),
    "MOVE_Z_Funding": ("MOVE_Z",),
}

BUSINESS_FIELDS: dict[str, tuple[str, ...]] = {
    "BusinessCyclePhase": ("BusinessCycleState",),
    "BusinessCycleLevel": ("BusinessCycleLevel",),
    "BusinessCycleMomentum": ("BusinessCycleMomentum",),
    "BusinessCycleDirection": ("BusinessCycleDirection",),
    "BusinessCycleTransitionZone": ("BusinessCycleTransitionZone",),
    "BusinessCycleConfidence": ("BusinessCycleConfidence",),
    "SurveyScore": ("SurveyScore",),
    "LaborScore": ("LaborScore",),
    "ProductionScore": ("ProductionScore",),
    "DemandIncomeScore": ("DemandIncomeScore",),
    "LaborCycleState": ("LaborCycleState",),
    "ProductivityExpansionFlag": ("ProductivityExpansionFlag",),
    "ISM_Manufacturing": ("NAPM",),
    "CFNAI": ("CFNAI",),
    "InitialClaims": ("ICSA",),
    "ContinuingClaims": ("CCSA",),
    "UnemploymentRate": ("UNRATE",),
    "Payrolls": ("PAYEMS",),
    "IndustrialProduction": ("INDPRO",),
    "RetailSales": ("RSAFS",),
    "RealPCE": ("PCEC96",),
    "RealPersonalIncomeExTransfers": ("W875RX1",),
    "ISM_Manufacturing_Z": ("ISM_Z",),
    "CFNAI_Z": ("CFNAI_Z",),
    "InitialClaims_Z": ("InitialClaims_Z_INV",),
    "ContinuingClaims_Z": ("ContinuingClaims_Z_INV",),
    "UnemploymentRate_Z": ("Unemployment_Z_INV",),
    "Payrolls_Z": ("Payrolls_Z",),
    "IndustrialProduction_Z": ("IndustrialProduction_Z",),
    "RetailSales_Z": ("RetailSales_Z",),
    "RealPCE_Z": ("RealPCE_Z",),
    "RealPersonalIncomeExTransfers_Z": ("RealPersonalIncome_Z",),
    "InflationDirection": ("InflationState",),
    "InflationLeadingScore": ("InflationDirectionScore",),
    "MarketPricingInflationScore": ("MarketPricingScore",),
    "ModelInflationScore": ("ModelImpliedInflationScore",),
    "SurveyInflationScore": ("SurveyInflationScore",),
    "RealizedInflationMomentum": ("RealizedInflationMomentum",),
    "InflationTransitionFlag": ("InflationTransitionZone",),
    "EconomyRegime": ("EconomyRegime",),
    "T5YIE": ("T5YIE",),
    "T10YIE": ("T10YIE",),
    "EXPINF1YR": ("EXPINF1YR",),
    "EXPINF5YR": ("EXPINF5YR",),
    "Michigan1Y": ("MICH",),
    "T5YIFR": ("T5YIFR",),
    "CorePCE": ("PCEPILFE",),
    "CoreCPI": ("CPILFESL",),
    "HeadlineCPI": ("CPIAUCSL",),
    "PPI": ("PPIACO",),
}

REQUIRED_SCHEMA_FIELDS = [
    "LiquidityForecastState", "GlobalLiquidityDirection", "LiquidityPressure", "PolicyResponse",
    "TreasuryFinancingPressure", "RatesPressure", "StructuralExtensionPercentile_PIT",
    "SMA200WExtensionPercentile_PIT", "PositioningVulnerability", "ReserveVulnerability",
    "CoreFCLevel", "CoreFCDirection", "CreditLevel", "CreditDirection", "VIXLevelRisk",
    "VIXMomentumRisk", "VIXTermStructureRisk", "RealizedVolatility_20D", "RealizedVolRisk",
    "FundingState", "FundingCore",
    "MoneyMarketStress", "ReservePressure", "PersistentFundingFlag", "CollateralStress",
    "FastTransitionRisk", "MacroTransitionRisk", "CurrentMarketRiskState", "BusinessCyclePhase",
    "BusinessCycleLevel", "BusinessCycleMomentum", "EconomyRegime",
]

CORE_COVERAGE_FIELDS = [
    "LiquidityPressure", "PolicyResponse", "RatesPressure", "StructuralExtensionPercentile_PIT",
    "SMA200WExtensionPercentile_PIT", "PositioningVulnerability", "CoreFCLevel", "CreditLevel",
    "VIXLevelRisk", "FundingCore", "ReservePressure", "FastTransitionRisk", "MacroTransitionRisk",
    "BusinessCycleLevel", "InflationLeadingScore",
]


def completed_week_end(value: str | pd.Timestamp | None = None) -> pd.Timestamp:
    now = pd.Timestamp.now(tz="UTC").tz_localize(None).normalize()
    requested = pd.Timestamp(value).tz_localize(None).normalize() if value is not None else now
    requested = min(requested, now)
    end = requested - pd.Timedelta(days=int((requested.weekday() - 4) % 7))
    if requested == now and now.weekday() == 4:
        end -= pd.Timedelta(days=7)
    return end


def build_financial_fragility_validation_workbook(
    *,
    market_snapshot: Any,
    liquidity_regime: pd.DataFrame,
    forecast_frame: pd.DataFrame,
    business_snapshot: Any,
    macro_snapshot: Any,
    rates_snapshot: Any,
    funding_snapshot: Any,
    treasury_snapshot: Any,
    transition_snapshot: dict[str, Any] | None = None,
    transition_history: pd.DataFrame | None = None,
    wresbal_history: pd.DataFrame | None = None,
    start_date: str | pd.Timestamp = DEFAULT_START_DATE,
    end_date: str | pd.Timestamp | None = None,
) -> tuple[bytes, str]:
    dataset, dictionary, export_metadata = build_financial_fragility_validation_dataset(
        market_snapshot=market_snapshot,
        liquidity_regime=liquidity_regime,
        forecast_frame=forecast_frame,
        business_snapshot=business_snapshot,
        macro_snapshot=macro_snapshot,
        rates_snapshot=rates_snapshot,
        funding_snapshot=funding_snapshot,
        treasury_snapshot=treasury_snapshot,
        transition_snapshot=transition_snapshot,
        transition_history=transition_history,
        start_date=start_date,
        end_date=end_date,
    )
    prepared_wresbal = prepare_wresbal_history(wresbal_history, treasury_snapshot)
    if not prepared_wresbal.empty:
        export_metadata = pd.concat(
            [
                export_metadata,
                pd.DataFrame(
                    [
                        ("WRESBALHistoryRows", len(prepared_wresbal)),
                        ("WRESBALHistoryStart", prepared_wresbal["ObservationDate"].min().strftime("%Y-%m-%d")),
                        ("WRESBALHistoryEnd", prepared_wresbal["ObservationDate"].max().strftime("%Y-%m-%d")),
                    ],
                    columns=["Field", "Value"],
                ),
            ],
            ignore_index=True,
        )
    payload = write_financial_fragility_workbook(
        dataset, dictionary, export_metadata, wresbal_history=prepared_wresbal
    )
    validate_workbook_payload(payload, dataset)
    return payload, EXPORT_FILENAME


def build_financial_fragility_validation_dataset(
    *,
    market_snapshot: Any,
    liquidity_regime: pd.DataFrame,
    forecast_frame: pd.DataFrame,
    business_snapshot: Any,
    macro_snapshot: Any,
    rates_snapshot: Any,
    funding_snapshot: Any,
    treasury_snapshot: Any,
    transition_snapshot: dict[str, Any] | None = None,
    transition_history: pd.DataFrame | None = None,
    start_date: str | pd.Timestamp = DEFAULT_START_DATE,
    end_date: str | pd.Timestamp | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    start = pd.Timestamp(start_date).tz_localize(None).normalize()
    end = completed_week_end(end_date)
    if start > end:
        raise ValueError("Export start date is after the latest completed weekly observation")

    dataset = pd.DataFrame({"Date": pd.date_range(start, end, freq="W-FRI")})
    metadata = [SeriesMeta(
        column="Date", layer="METADATA", category="Dataset",
        description="Completed weekly observation date", unit="date",
        transformation="W-FRI completed-week calendar", source="Internal",
        original_frequency="weekly", model_usage="Primary key",
        point_in_time_safe="Yes", forward_filled=False,
    )]

    market_history = getattr(market_snapshot, "history", pd.DataFrame())
    market_daily = getattr(market_snapshot, "daily", pd.DataFrame())
    _merge_fields(dataset, market_history, "Date", MARKET_FIELDS, metadata, "Market Cycle")
    add_outcome_targets(dataset, metadata, market_daily)

    _merge_fields(dataset, liquidity_regime, "date", GLOBAL_LIQUIDITY_FIELDS, metadata, "Global Liquidity")
    if "GlobalLiquidityLevel" in dataset:
        score = pd.to_numeric(dataset["GlobalLiquidityLevel"], errors="coerce")
        dataset["GlobalLiquidityROC"] = score.diff(4)
        dataset["GlobalLiquidityROC_13W"] = score.diff(13)
        dataset["GlobalLiquidityROC_26W"] = score.diff(26)
        for column, weeks in (("GlobalLiquidityROC", 4), ("GlobalLiquidityROC_13W", 13), ("GlobalLiquidityROC_26W", 26)):
            _append_meta(metadata, dataset, column, "Global Liquidity", f"Change in production liquidity score over {weeks} weeks", "Point-in-time difference")

    _merge_fields(
        dataset,
        transition_history if transition_history is not None else pd.DataFrame(),
        "Date",
        TRANSITION_FIELDS,
        metadata,
        "Market Transition",
    )
    add_transition_snapshot(dataset, transition_snapshot or {}, metadata)

    _merge_fields(dataset, forecast_frame, "Date", LIQUIDITY_FORECAST_FIELDS, metadata, "Liquidity Forecast")

    rates_history = getattr(rates_snapshot, "history", pd.DataFrame())
    _merge_fields(dataset, rates_history, "Date", RATES_FIELDS, metadata, "Rates & Financial Conditions")

    funding_history = getattr(funding_snapshot, "weekly", pd.DataFrame())
    _merge_fields(dataset, funding_history, "Date", FUNDING_FIELDS, metadata, "Funding Conditions")

    treasury_history = getattr(treasury_snapshot, "weekly", pd.DataFrame())
    if treasury_history is None or treasury_history.empty:
        treasury_history = getattr(treasury_snapshot, "history", pd.DataFrame())
    _merge_fields(dataset, treasury_history, "Date", TREASURY_FIELDS, metadata, "Treasury & Fiscal")

    business_history = getattr(business_snapshot, "history", pd.DataFrame())
    _merge_fields(dataset, business_history, "date", BUSINESS_FIELDS, metadata, "Business Cycle / Inflation")
    surprise_history = getattr(macro_snapshot, "history", pd.DataFrame())
    _merge_fields(dataset, surprise_history, "Date", MACRO_SURPRISE_FIELDS, metadata, "Macro Surprises")

    _add_aliases(dataset, BASE_ALIASES, metadata)
    add_cross_cycle_state(dataset, metadata)

    _add_required_missing_columns(dataset, metadata)
    add_data_quality_fields(dataset, metadata)
    dataset = _order_columns(dataset)
    validate_fragility_dataset(dataset)
    dictionary = build_data_dictionary(dataset, metadata)
    export_metadata = build_export_metadata(dataset)
    return dataset, dictionary, export_metadata


def _normalize_weekly(dataset: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp) -> pd.DataFrame:
    frame = dataset.copy()
    frame["Date"] = pd.to_datetime(frame["Date"], errors="coerce").dt.tz_localize(None).dt.normalize()
    frame = frame.dropna(subset=["Date"]).loc[lambda x: x["Date"].between(start, end)]
    return frame.sort_values("Date").drop_duplicates("Date", keep="last").reset_index(drop=True)


def _merge_fields(
    dataset: pd.DataFrame,
    source: pd.DataFrame,
    date_column: str,
    mapping: dict[str, tuple[str, ...]],
    metadata: list[SeriesMeta],
    module: str,
) -> None:
    if source is None or source.empty or date_column not in source:
        for output in mapping:
            _ensure_missing(dataset, output, metadata, module, "Production history unavailable")
        return
    frame = source.copy()
    frame[date_column] = pd.to_datetime(frame[date_column], errors="coerce").dt.tz_localize(None).dt.normalize()
    frame = frame.dropna(subset=[date_column]).sort_values(date_column).drop_duplicates(date_column, keep="last")
    aligned = frame.set_index(date_column).reindex(pd.DatetimeIndex(dataset["Date"]))
    for output, candidates in mapping.items():
        source_column = next((column for column in candidates if column in aligned), None)
        if source_column is None:
            _ensure_missing(dataset, output, metadata, module, f"Not exposed by current {module} production history")
            continue
        dataset[output] = aligned[source_column].to_numpy()
        _append_meta(metadata, dataset, output, module, f"Exact production field: {source_column}", "Production W-FRI history")


def _add_aliases(dataset: pd.DataFrame, mapping: dict[str, tuple[str, ...]], metadata: list[SeriesMeta]) -> None:
    for output, candidates in mapping.items():
        if output in dataset and dataset[output].notna().any():
            continue
        source = next((column for column in candidates if column in dataset), None)
        if source is None:
            continue
        dataset[output] = dataset[source]
        _append_meta(metadata, dataset, output, "Production Alias", f"Alias of production field {source}", "No recalculation")


def add_transition_snapshot(
    dataset: pd.DataFrame,
    snapshot: dict[str, Any],
    metadata: list[SeriesMeta],
) -> None:
    for output, candidates in TRANSITION_FIELDS.items():
        value = next((snapshot.get(column) for column in candidates if snapshot.get(column) is not None), np.nan)
        if output not in dataset:
            dataset[output] = pd.Series(pd.NA, index=dataset.index, dtype="object") if isinstance(value, str) else np.nan
        latest_missing = output not in dataset or pd.isna(dataset[output].iloc[-1])
        if len(dataset) and latest_missing and not (isinstance(value, float) and np.isnan(value)):
            dataset.loc[dataset.index[-1], output] = value
        _append_meta(
            metadata,
            dataset,
            output,
            "Market Transition",
            f"Current loaded snapshot field: {next((column for column in candidates if column in snapshot), candidates[0])}",
            "Historical production series with latest loaded snapshot fallback",
        )


def prepare_wresbal_history(
    wresbal_history: pd.DataFrame | None,
    treasury_snapshot: Any,
) -> pd.DataFrame:
    columns = ["ObservationDate", "AvailableDate", "WRESBAL_USD_Millions", "WRESBAL_USD_Bn"]
    if wresbal_history is not None and not wresbal_history.empty:
        frame = wresbal_history.copy()
    else:
        liquidity = getattr(treasury_snapshot, "liquidity", pd.DataFrame())
        if liquidity is None or liquidity.empty or not {"Date", "WRESBAL"}.issubset(liquidity.columns):
            return pd.DataFrame(columns=columns)
        frame = liquidity[["Date", "WRESBAL"]].rename(columns={"Date": "ObservationDate"})
        frame["AvailableDate"] = frame["ObservationDate"]
        frame["WRESBAL_USD_Bn"] = pd.to_numeric(frame["WRESBAL"], errors="coerce")
        frame["WRESBAL_USD_Millions"] = frame["WRESBAL_USD_Bn"] * 1000.0
    for date_column in ("ObservationDate", "AvailableDate"):
        frame[date_column] = pd.to_datetime(frame[date_column], errors="coerce").dt.tz_localize(None)
    if "WRESBAL_USD_Millions" not in frame and "Value" in frame:
        raw = pd.to_numeric(frame["Value"], errors="coerce")
        frame["WRESBAL_USD_Millions"] = raw.where(raw.ge(100_000), raw * 1000)
    if "WRESBAL_USD_Bn" not in frame:
        frame["WRESBAL_USD_Bn"] = pd.to_numeric(frame["WRESBAL_USD_Millions"], errors="coerce") / 1000.0
    return (
        frame[columns]
        .dropna(subset=["ObservationDate", "WRESBAL_USD_Millions"])
        .sort_values(["ObservationDate", "AvailableDate"])
        .drop_duplicates("ObservationDate", keep="first")
        .reset_index(drop=True)
    )


def _ensure_missing(
    dataset: pd.DataFrame,
    column: str,
    metadata: list[SeriesMeta],
    module: str,
    note: str,
) -> None:
    if column not in dataset:
        dataset[column] = np.nan
        _append_meta(metadata, dataset, column, module, note, "Unavailable; values intentionally left missing")


def _add_required_missing_columns(dataset: pd.DataFrame, metadata: list[SeriesMeta]) -> None:
    for field in REQUIRED_SCHEMA_FIELDS:
        _ensure_missing(dataset, field, metadata, "Requested Core Schema", "Production field unavailable")


def _append_meta(
    metadata: list[SeriesMeta],
    dataset: pd.DataFrame,
    column: str,
    module: str,
    description: str,
    transformation: str,
    *,
    target: bool = False,
) -> None:
    if any(item.column == column for item in metadata):
        return
    values = dataset[column]
    if pd.api.types.is_bool_dtype(values.dtype):
        unit = "boolean"
    elif pd.api.types.is_datetime64_any_dtype(values.dtype):
        unit = "date"
    elif values.dtype == "object":
        unit = "state / text"
    else:
        unit = "numeric"
    metadata.append(SeriesMeta(
        column=column,
        layer="VALIDATION_TARGET" if target else "MODEL_OUTPUT",
        category=module,
        description=description,
        unit=unit,
        transformation=transformation,
        source="Research export" if target else f"Production {module}",
        original_frequency="weekly",
        model_usage="Validation target only - never a model feature" if target else "Financial Fragility validation input",
        point_in_time_safe="Contains future data - target only" if target else "Yes - production point-in-time convention",
        forward_filled=False,
        notes="FOR RESEARCH / VALIDATION ONLY" if target else "",
    ))


def add_outcome_targets(
    dataset: pd.DataFrame,
    metadata: list[SeriesMeta],
    daily_prices: pd.DataFrame | None = None,
) -> None:
    close = pd.to_numeric(dataset.get("SPX_Close"), errors="coerce")
    if close.isna().all() and "SPY_Close" in dataset:
        close = pd.to_numeric(dataset["SPY_Close"], errors="coerce")
    for label, weeks in HORIZONS.items():
        backward = close.pct_change(weeks, fill_method=None)
        forward_return = close.shift(-weeks) / close - 1.0
        from_current, peak_to_trough = _forward_drawdowns(close, weeks)
        if daily_prices is not None and not daily_prices.empty:
            daily_from_current, daily_peak_to_trough = _daily_forward_drawdowns(
                pd.DatetimeIndex(dataset["Date"]), daily_prices, weeks
            )
            from_current = daily_from_current
            peak_to_trough = daily_peak_to_trough
        columns = {
            f"SPX_Return_{label}": backward,
            f"SPX_Forward_Return_{label}": forward_return,
            f"SPX_Forward_MaxDD_{label}": peak_to_trough,
            f"SPX_Forward_MaxDD_FromCurrent_{label}": from_current,
            f"SPX_Forward_PeakToTroughMaxDD_{label}": peak_to_trough,
        }
        for column, values in columns.items():
            dataset[column] = values
            _append_meta(
                metadata,
                dataset,
                column,
                "Market Outcomes",
                "Historical return" if column.startswith("SPX_Return_") else "Forward validation outcome",
                f"Computed from weekly SPX close over {weeks} completed weeks",
                target=not column.startswith("SPX_Return_"),
            )
    dataset["SPX_Return_1W"] = close.pct_change(1, fill_method=None)
    _append_meta(metadata, dataset, "SPX_Return_1W", "Market Outcomes", "Historical one-week SPX return", "Close / prior close - 1")

    thresholds = {
        "1M": (15, 20),
        "3M": (15, 20, 25),
        "6M": (15, 20, 25, 35),
        "12M": (15, 20, 25, 35, 45),
    }
    for label, levels in thresholds.items():
        drawdown = pd.to_numeric(dataset[f"SPX_Forward_MaxDD_{label}"], errors="coerce")
        for level in levels:
            column = f"ForwardDD{level}_{label}"
            dataset[column] = drawdown.le(-level / 100.0).where(drawdown.notna()).astype("boolean")
            _append_meta(
                metadata,
                dataset,
                column,
                "Market Outcomes",
                f"Forward peak-to-trough drawdown reached {level}%",
                f"1 when SPX_Forward_MaxDD_{label} <= -{level / 100.0:.2f}",
                target=True,
            )


def _forward_drawdowns(close: pd.Series, weeks: int) -> tuple[pd.Series, pd.Series]:
    values = pd.to_numeric(close, errors="coerce").to_numpy(dtype=float)
    from_current = np.full(len(values), np.nan)
    peak_to_trough = np.full(len(values), np.nan)
    for index in range(len(values)):
        end = index + weeks
        if end >= len(values) or not np.isfinite(values[index]):
            continue
        window = values[index : end + 1]
        if not np.isfinite(window).all():
            continue
        from_current[index] = np.min(window[1:] / values[index] - 1.0)
        running_peak = np.maximum.accumulate(window)
        peak_to_trough[index] = np.min(window / running_peak - 1.0)
    return pd.Series(from_current, index=close.index), pd.Series(peak_to_trough, index=close.index)


def _daily_forward_drawdowns(
    weekly_dates: pd.DatetimeIndex,
    daily_prices: pd.DataFrame,
    weeks: int,
) -> tuple[pd.Series, pd.Series]:
    daily = daily_prices.copy()
    if "Date" in daily:
        daily.index = pd.to_datetime(daily["Date"], errors="coerce")
    else:
        daily.index = pd.to_datetime(daily.index, errors="coerce")
    close_column = "SPX_Close" if "SPX_Close" in daily else "SPY_RawClose" if "SPY_RawClose" in daily else "SPY_Close"
    if close_column not in daily:
        return pd.Series(np.nan, index=range(len(weekly_dates))), pd.Series(np.nan, index=range(len(weekly_dates)))
    daily = daily.loc[daily.index.notna()].sort_index()
    daily = daily.loc[~daily.index.duplicated(keep="last")]
    close = pd.to_numeric(daily[close_column], errors="coerce")
    low = pd.to_numeric(daily.get("SPX_Low", close), errors="coerce")
    from_current = []
    peak_to_trough = []
    last_daily_date = daily.index.max() if not daily.empty else pd.NaT
    for anchor in pd.to_datetime(weekly_dates):
        end = anchor + pd.Timedelta(weeks=weeks)
        if pd.isna(last_daily_date) or end > last_daily_date:
            from_current.append(np.nan)
            peak_to_trough.append(np.nan)
            continue
        window_close = close.loc[(close.index > anchor) & (close.index <= end)].dropna()
        window_low = low.reindex(window_close.index).dropna()
        anchor_values = close.loc[close.index <= anchor].dropna()
        if anchor_values.empty or window_close.empty or window_low.empty:
            from_current.append(np.nan)
            peak_to_trough.append(np.nan)
            continue
        anchor_close = float(anchor_values.iloc[-1])
        from_current.append(float((window_low / anchor_close - 1.0).min()))
        combined = pd.concat([pd.Series([anchor_close], index=[anchor]), window_close])
        running_peak = combined.cummax()
        peak_to_trough.append(float((combined / running_peak - 1.0).min()))
    return pd.Series(from_current, index=range(len(weekly_dates))), pd.Series(peak_to_trough, index=range(len(weekly_dates)))


def add_cross_cycle_state(dataset: pd.DataFrame, metadata: list[SeriesMeta]) -> None:
    required = ("LiquidityForecastState", "SPXMomentumCyclePhase", "BusinessCycleDirection")
    if not all(column in dataset for column in required):
        return
    dataset["CrossCycleState"] = [
        classify_cross_cycle(liquidity, market, economy)
        for liquidity, market, economy in zip(*(dataset[column] for column in required))
    ]
    _append_meta(
        metadata,
        dataset,
        "CrossCycleState",
        "Cross-Cycle",
        "Existing Global Dashboard sequence classification",
        "Liquidity Forecast State x SPX Momentum Cycle x Business Cycle Direction",
    )


def add_data_quality_fields(dataset: pd.DataFrame, metadata: list[SeriesMeta]) -> None:
    available_fields = [field for field in CORE_COVERAGE_FIELDS if field in dataset]
    available = dataset[available_fields].notna().sum(axis=1) if available_fields else pd.Series(0, index=dataset.index)
    total = len(CORE_COVERAGE_FIELDS)
    dataset["NumberOfAvailableCoreInputs"] = available.astype(int)
    dataset["NumberOfMissingCoreInputs"] = (total - available).astype(int)
    dataset["DataCoveragePct"] = available / total * 100.0
    dataset["PartialDataFlag"] = available.lt(total)
    dataset["StaleDataFlag"] = False
    for column, description in {
        "NumberOfAvailableCoreInputs": "Count of available core validation inputs",
        "NumberOfMissingCoreInputs": "Count of missing core validation inputs",
        "DataCoveragePct": "Available core inputs divided by requested core inputs",
        "PartialDataFlag": "TRUE when at least one core validation input is missing",
        "StaleDataFlag": "Historical rows use completed production observations; staleness is not inferred from missing values",
    }.items():
        _append_meta(metadata, dataset, column, "Data Quality", description, "Row-level coverage calculation")


def _order_columns(dataset: pd.DataFrame) -> pd.DataFrame:
    outcomes = [column for column in dataset if column.startswith(("SPX_Return_", "SPX_Forward_", "ForwardDD"))]
    quality = ["DataCoveragePct", "NumberOfAvailableCoreInputs", "NumberOfMissingCoreInputs", "PartialDataFlag", "StaleDataFlag"]
    priority = ["Date", "SPX_Close", "SPY_Close"] + REQUIRED_SCHEMA_FIELDS + quality
    ordered = []
    for column in priority + [column for column in dataset if column not in outcomes] + outcomes:
        if column in dataset and column not in ordered:
            ordered.append(column)
    return dataset[ordered]


def build_data_dictionary(dataset: pd.DataFrame, metadata: Iterable[SeriesMeta]) -> pd.DataFrame:
    by_column = {item.column: item for item in metadata}
    rows = []
    dates = pd.to_datetime(dataset["Date"], errors="coerce")
    for column in dataset.columns:
        meta = by_column.get(column)
        valid = dataset[column].notna()
        earliest = dates.loc[valid].min() if valid.any() else pd.NaT
        rows.append({
            "FieldName": column,
            "Module": meta.category if meta else "Unclassified production field",
            "Description": meta.description if meta else column,
            "Formula / Source": f"{meta.transformation}; {meta.source}" if meta else "Existing production export field",
            "Frequency": meta.original_frequency if meta else "weekly",
            "Unit": meta.unit if meta else _infer_unit(column, dataset[column]),
            "HigherMeans": _higher_means(column),
            "PITSafe": meta.point_in_time_safe if meta else "Production timing convention",
            "EarliestValidDate": earliest.strftime("%Y-%m-%d") if pd.notna(earliest) else None,
            "Notes": meta.notes if meta else "",
        })
    return pd.DataFrame(rows)


def _infer_unit(column: str, values: pd.Series) -> str:
    if pd.api.types.is_bool_dtype(values.dtype):
        return "boolean"
    if pd.api.types.is_datetime64_any_dtype(values.dtype) or column.endswith(("Date", "LastUpdated")):
        return "date"
    if values.dtype == "object":
        return "state / text"
    if "Percentile" in column or column.endswith("Pct") or "Return" in column or "MaxDD" in column:
        return "decimal fraction unless source definition states 0-100"
    return "numeric"


def _higher_means(column: str) -> str:
    lower = column.lower()
    if any(token in lower for token in ("risk", "stress", "pressure", "vulnerability", "drawdown")):
        return "Higher stress / vulnerability; drawdowns are more severe when more negative"
    if any(token in lower for token in ("liquidity", "growth", "return", "impulse")):
        return "Higher value generally indicates stronger impulse / outcome; see production definition"
    return "See production definition"


def build_export_metadata(dataset: pd.DataFrame) -> pd.DataFrame:
    versions = {
        "GLOBAL_LIQUIDITY": "GL_Recomputed_Current_Model",
        "LIQUIDITY_FORECAST": FORECAST_VERSION,
        "BREADTH": BREADTH_VERSION,
        "MARKET_CYCLE": MARKET_CYCLE_MODEL_VERSION,
        "BUSINESS_CYCLE": BUSINESS_CYCLE_MODEL_VERSION,
        "INFLATION_LAYER": INFLATION_LAYER_MODEL_VERSION,
        "ECONOMY_REGIME": ECONOMY_REGIME_MODEL_VERSION,
        "MACRO_SURPRISES": MACRO_SURPRISES_MODEL_VERSION,
        "RATES_FINANCIAL_CONDITIONS": RATES_FC_MODEL_VERSION,
        "FUNDING_CONDITIONS": FUNDING_MODEL_VERSION,
        "TREASURY_FISCAL": TREASURY_FISCAL_MODEL_VERSION,
    }
    rows = [
        ("ExportDate", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")),
        ("AppVersion / GitCommit", os.getenv("GIT_COMMIT", "UNAVAILABLE")),
        ("EarliestDate", pd.Timestamp(dataset["Date"].min()).strftime("%Y-%m-%d")),
        ("LatestDate", pd.Timestamp(dataset["Date"].max()).strftime("%Y-%m-%d")),
        ("Frequency", "WEEKLY / completed W-FRI"),
        ("NumberOfRows", len(dataset)),
        ("NumberOfColumns", len(dataset.columns)),
        ("ForwardOutcomeColumnsContainFutureData", True),
        ("ForwardOutcomeWarning", "FOR RESEARCH / VALIDATION ONLY - DO NOT USE AS MODEL FEATURES."),
        ("FinancialFragilityPrototypeIncluded", False),
        ("FinancialFragilityPrototypeNote", "No production FINANCIAL_FRAGILITY_V1 fields existed at export time; no placeholders were created."),
    ]
    rows.extend((f"ModelVersion_{name}", version) for name, version in versions.items())
    return pd.DataFrame(rows, columns=["Field", "Value"])


def write_financial_fragility_workbook(
    dataset: pd.DataFrame,
    dictionary: pd.DataFrame,
    export_metadata: pd.DataFrame,
    wresbal_history: pd.DataFrame | None = None,
) -> bytes:
    output = BytesIO()
    workbook = Workbook(write_only=True)
    write_dataframe_sheet(workbook, "FRAGILITY_VALIDATION_DATA", dataset, freeze_first_column=True)
    write_dataframe_sheet(workbook, "DATA_DICTIONARY", dictionary)
    write_dataframe_sheet(workbook, "EXPORT_METADATA", export_metadata)
    if wresbal_history is not None and not wresbal_history.empty:
        write_dataframe_sheet(workbook, "WRESBAL_HISTORY", wresbal_history)
    workbook.save(output)
    return output.getvalue()


def validate_fragility_dataset(dataset: pd.DataFrame) -> None:
    if dataset.empty:
        raise ValueError("FRAGILITY_VALIDATION_DATA is empty")
    dates = pd.to_datetime(dataset["Date"], errors="coerce")
    if dates.isna().any() or dates.duplicated().any() or not dates.is_monotonic_increasing:
        raise ValueError("Weekly Date must be valid, unique and ascending")
    if not dates.dt.weekday.eq(4).all():
        raise ValueError("Weekly Date must use the completed Friday anchor")
    missing = [field for field in ["SPX_Close", *REQUIRED_SCHEMA_FIELDS] if field not in dataset]
    if missing:
        raise ValueError(f"Required validation fields missing: {missing}")
    target_columns = [column for column in dataset if column.startswith(("SPX_Forward_", "ForwardDD"))]
    if not target_columns:
        raise ValueError("Forward outcome targets are missing")


def validate_workbook_payload(payload: bytes, dataset: pd.DataFrame) -> None:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    required = {"FRAGILITY_VALIDATION_DATA", "DATA_DICTIONARY", "EXPORT_METADATA"}
    if not required.issubset(workbook.sheetnames):
        raise ValueError("Workbook is missing required validation sheets")
    sheet = workbook["FRAGILITY_VALIDATION_DATA"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, ())
    row_count = sum(1 for _ in rows)
    if row_count != len(dataset) or len(header) != len(dataset.columns):
        raise ValueError("Workbook data dimensions do not match the validated dataset")
    workbook.close()
