from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


WGC_FIELDS = {
    "Total Supply": "total_supply_tonnes",
    "Total Mine Supply": "total_mine_supply_tonnes",
    "Recycled Gold": "recycled_gold_tonnes",
    "Jewellery Fabrication": "jewellery_fabrication_tonnes",
    "Technology": "technology_tonnes",
    "Investment": "investment_tonnes",
    "Central Bank and Other Institutions": "central_banks_tonnes",
    "OTC and other": "otc_and_other_tonnes",
    "LBMA Gold Price (US$/oz)": "lbma_gold_price_usd_oz",
}


def import_wgc_workbook(source: str | Path, output_dir: str | Path) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    """Normalize the WGC presentation workbook into stable annual/quarterly datasets."""
    source_path = Path(source)
    output_path = Path(output_dir)
    workbook = load_workbook(source_path, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    row_by_label: dict[str, int] = {}
    for row in range(1, sheet.max_row + 1):
        label = sheet.cell(row, 2).value
        if label is not None:
            row_by_label[str(label).strip()] = row
    missing = sorted(set(WGC_FIELDS) - set(row_by_label))
    if missing:
        raise ValueError(f"WGC workbook is missing required rows: {', '.join(missing)}")

    annual_rows: list[dict[str, Any]] = []
    quarterly_rows: list[dict[str, Any]] = []
    for column in range(1, sheet.max_column + 1):
        header = sheet.cell(5, column).value
        if isinstance(header, (int, float)) and 1900 <= int(header) <= 2200:
            annual_rows.append(_record(sheet, column, str(int(header)), row_by_label))
            continue
        match = re.fullmatch(r"Q([1-4])'(\d{2})", str(header or "").strip())
        if match:
            quarter = int(match.group(1))
            year = 2000 + int(match.group(2))
            record = _record(sheet, column, f"{year}Q{quarter}", row_by_label)
            record.update({"year": year, "quarter": quarter})
            quarterly_rows.append(record)

    annual = pd.DataFrame(annual_rows).sort_values("period").reset_index(drop=True)
    quarterly = pd.DataFrame(quarterly_rows).sort_values(["year", "quarter"]).reset_index(drop=True)
    annual["year"] = pd.to_numeric(annual["period"], errors="raise").astype(int)
    value_columns = list(WGC_FIELDS.values())
    annual["is_complete"] = annual[value_columns].notna().all(axis=1)
    quarterly["is_published"] = quarterly[value_columns].notna().any(axis=1)

    as_of = None
    source_label = None
    for row in range(1, sheet.max_row + 1):
        label = str(sheet.cell(row, 2).value or "").strip()
        if label.lower().startswith("data as of"):
            as_of = label.removeprefix("Data as of ")
        elif label.lower().startswith("source:"):
            source_label = label.removeprefix("Source: ")

    metadata = {
        "dataset": "World Gold Council gold supply and demand presentation",
        "source_file": source_path.name,
        "source_sha256": hashlib.sha256(source_path.read_bytes()).hexdigest(),
        "data_as_of": as_of,
        "source": source_label,
        "annual_first_year": int(annual["year"].min()),
        "annual_last_year": int(annual["year"].max()),
        "latest_published_quarter": str(quarterly.loc[quarterly["is_published"], "period"].iloc[-1]),
        "runtime_dependency_on_source_file": False,
    }

    output_path.mkdir(parents=True, exist_ok=True)
    annual.to_csv(output_path / "wgc_gold_balance_annual.csv", index=False, float_format="%.10f")
    quarterly.to_csv(output_path / "wgc_gold_balance_quarterly.csv", index=False, float_format="%.10f")
    (output_path / "wgc_gold_balance_metadata.json").write_text(
        json.dumps(metadata, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return annual, quarterly, metadata


def _record(sheet: Any, column: int, period: str, row_by_label: dict[str, int]) -> dict[str, Any]:
    record: dict[str, Any] = {"period": period}
    for source_label, output_name in WGC_FIELDS.items():
        record[output_name] = pd.to_numeric(sheet.cell(row_by_label[source_label], column).value, errors="coerce")
    return record


def main() -> None:
    parser = argparse.ArgumentParser(description="Normalize a WGC Gold Balance workbook for Gold Regime.")
    parser.add_argument("source", help="Path to the WGC .xlsx report")
    parser.add_argument(
        "--output-dir",
        default=str(Path("data") / "gold_regime"),
        help="Repository directory for normalized CSV/JSON files",
    )
    args = parser.parse_args()
    annual, quarterly, metadata = import_wgc_workbook(args.source, args.output_dir)
    print(
        f"Imported {len(annual)} annual rows and {len(quarterly)} quarterly rows; "
        f"latest published quarter: {metadata['latest_published_quarter']}"
    )


if __name__ == "__main__":
    main()
