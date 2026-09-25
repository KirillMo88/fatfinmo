from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from financial_fragility_export import (
    _forward_drawdowns,
    add_data_quality_fields,
    add_outcome_targets,
    build_financial_fragility_validation_dataset,
    build_data_dictionary,
    build_export_metadata,
    completed_week_end,
    prepare_wresbal_history,
    validate_fragility_dataset,
    write_financial_fragility_workbook,
)


def test_forward_outcomes_are_targets_and_peak_to_trough_is_true_drawdown():
    close = pd.Series([100.0, 120.0, 90.0, 110.0, 130.0])
    from_current, peak_to_trough = _forward_drawdowns(close, 3)
    assert from_current.iloc[0] == pytest.approx(-0.10)
    assert peak_to_trough.iloc[0] == pytest.approx(-0.25)


def test_validation_dataset_and_workbook_have_required_shape():
    dates = pd.date_range("2024-01-05", periods=60, freq="W-FRI")
    frame = pd.DataFrame({"Date": dates, "SPX_Close": np.linspace(4000, 5000, len(dates))})
    required = [
        "LiquidityForecastState", "GlobalLiquidityDirection", "LiquidityPressure", "PolicyResponse",
        "TreasuryFinancingPressure", "RatesPressure", "StructuralExtensionPercentile_PIT",
        "SMA200WExtensionPercentile_PIT", "PositioningVulnerability", "ReserveVulnerability",
        "CoreFCLevel", "CoreFCDirection", "CreditLevel", "CreditDirection", "VIXLevelRisk",
        "VIXMomentumRisk", "VIXTermStructureRisk", "RealizedVolatility_20D", "RealizedVolRisk",
        "FundingState", "FundingCore",
        "MoneyMarketStress", "ReservePressure", "PersistentFundingFlag", "CollateralStress",
        "FastTransitionRisk", "MacroTransitionRisk", "CurrentMarketRiskState", "BusinessCyclePhase",
        "BusinessCycleLevel", "BusinessCycleMomentum", "EconomyRegime",
    ]
    for column in required:
        frame[column] = "NORMAL" if column.endswith(("State", "Direction", "Phase")) else 1.0
    metadata = []
    add_outcome_targets(frame, metadata)
    add_data_quality_fields(frame, metadata)
    validate_fragility_dataset(frame)
    dictionary = build_data_dictionary(frame, metadata)
    export_metadata = build_export_metadata(frame)
    payload = write_financial_fragility_workbook(frame, dictionary, export_metadata)
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    assert workbook.sheetnames == ["FRAGILITY_VALIDATION_DATA", "DATA_DICTIONARY", "EXPORT_METADATA"]
    rows = list(workbook["FRAGILITY_VALIDATION_DATA"].iter_rows(values_only=True))
    assert len(rows) == len(frame) + 1
    assert set(frame.columns) == set(dictionary["FieldName"])
    warning = export_metadata.set_index("Field").loc["ForwardOutcomeWarning", "Value"]
    assert "DO NOT USE AS MODEL FEATURES" in warning


def test_completed_week_end_never_returns_non_friday():
    assert completed_week_end("2024-01-10") == pd.Timestamp("2024-01-05")


def test_export_assembles_loaded_snapshots_without_recalculation():
    dates = pd.date_range("2024-01-05", periods=60, freq="W-FRI")
    market = SimpleNamespace(
        history=pd.DataFrame({"Date": dates, "SPX_Close": np.linspace(4000, 5000, 60)}),
        daily=pd.DataFrame(),
    )
    empty_date = SimpleNamespace(history=pd.DataFrame(), weekly=pd.DataFrame())
    dataset, dictionary, metadata = build_financial_fragility_validation_dataset(
        market_snapshot=market,
        liquidity_regime=pd.DataFrame(),
        forecast_frame=pd.DataFrame(),
        business_snapshot=SimpleNamespace(history=pd.DataFrame()),
        macro_snapshot=SimpleNamespace(history=pd.DataFrame()),
        rates_snapshot=empty_date,
        funding_snapshot=empty_date,
        treasury_snapshot=empty_date,
        transition_snapshot={"Fast_Transition_Risk": 17.0, "Macro_Transition_Risk": 42.0},
        start_date=dates.min(),
        end_date=dates.max(),
    )
    assert len(dataset) == 60
    assert dataset["FastTransitionRisk"].iloc[-1] == 17.0
    assert dataset["FastTransitionRisk"].iloc[:-1].isna().all()
    assert set(dataset.columns) == set(dictionary["FieldName"])
    assert metadata.set_index("Field").loc["NumberOfRows", "Value"] == 60


def test_requested_fragility_histories_are_merged_for_every_available_week():
    dates = pd.date_range("2024-01-05", periods=60, freq="W-FRI")
    market_history = pd.DataFrame({
        "Date": dates,
        "SPX_Close": np.linspace(4000, 5000, len(dates)),
        "RealizedVol20D": np.linspace(0.10, 0.25, len(dates)),
        "RealizedVolRisk": np.linspace(10, 90, len(dates)),
    })
    transition_history = pd.DataFrame({
        "Date": dates,
        "Fast_Transition_Risk": np.linspace(5, 55, len(dates)),
        "Macro_Transition_Risk": np.linspace(10, 60, len(dates)),
        "PositioningRisk": np.linspace(20, 80, len(dates)),
        "PositioningState": "NORMAL",
    })
    treasury_history = pd.DataFrame({
        "Date": dates,
        "TreasuryFinancingPressure": "MODERATE",
        "AbsorptionTightness": np.linspace(25, 85, len(dates)),
        "TreasuryLiquidityImpulse": np.linspace(-1, 1, len(dates)),
        "FiscalImpulse": np.linspace(-0.5, 0.5, len(dates)),
        "WRESBAL": np.linspace(3000, 3400, len(dates)),
    })
    empty = SimpleNamespace(history=pd.DataFrame(), weekly=pd.DataFrame())
    dataset, _, _ = build_financial_fragility_validation_dataset(
        market_snapshot=SimpleNamespace(history=market_history, daily=pd.DataFrame()),
        liquidity_regime=pd.DataFrame(),
        forecast_frame=pd.DataFrame(),
        business_snapshot=SimpleNamespace(history=pd.DataFrame()),
        macro_snapshot=SimpleNamespace(history=pd.DataFrame()),
        rates_snapshot=empty,
        funding_snapshot=empty,
        treasury_snapshot=SimpleNamespace(weekly=treasury_history),
        transition_snapshot={},
        transition_history=transition_history,
        start_date=dates.min(),
        end_date=dates.max(),
    )
    expected = [
        "TreasuryFinancingPressure", "FinancingPressureScore", "PositioningVulnerability",
        "FastTransitionRisk", "MacroTransitionRisk", "RealizedVolatility_20D",
        "RealizedVolRisk", "WRESBAL", "TreasuryLiquidityImpulse", "FiscalGrowthImpulse",
    ]
    assert dataset[expected].notna().all().all()
    assert dataset["PositioningVulnerability"].iloc[0] == pytest.approx(20.0)
    assert dataset["FastTransitionRisk"].iloc[-1] == pytest.approx(55.0)
    assert dataset["FinancingPressureScore"].iloc[-1] == pytest.approx(85.0)


def test_full_wresbal_history_is_independent_of_sofr_window():
    observations = pd.date_range("2010-01-06", periods=4, freq="W-WED")
    raw = pd.DataFrame({
        "ObservationDate": observations,
        "AvailableDate": observations + pd.Timedelta(days=1),
        "WRESBAL_USD_Millions": [1_000_000, 1_010_000, 1_020_000, 1_030_000],
        "WRESBAL_USD_Bn": [1000, 1010, 1020, 1030],
    })
    prepared = prepare_wresbal_history(raw, SimpleNamespace(liquidity=pd.DataFrame()))
    assert prepared["ObservationDate"].min() == pd.Timestamp("2010-01-06")
    assert list(prepared.columns) == [
        "ObservationDate", "AvailableDate", "WRESBAL_USD_Millions", "WRESBAL_USD_Bn",
    ]
