from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from btc_cycle import (
    BTC_LIQUIDITY_CYCLE_MONTHS,
    BTC_PROJECTED_HALVING,
    BTC_RANGE_OPTIONS,
    BTC_MODULAR_CYCLE_START,
    BTC_MODULAR_HISTORICAL_MODULES,
    build_btc_modular_cycle_forecast,
    build_btc_cycle_history,
    build_btc_gold_ratio_history,
    merge_btc_mcp_weekly_history,
    btc_cycle_time_range,
    btc_cycle_export_xlsx,
    btc_cycle_validation,
    halving_cycle_position,
    halving_phase,
    liquidity_cycle_position,
    next_accumulation_pre_halving_start,
    point_in_time_percentile,
)
from btc_cycle_tab import (
    _build_btc_etf_flow_intensity_figure,
    _build_btc_gold_ratio_figure,
    _build_btc_price_halving_figure,
    _build_macro_score_figure,
    _build_structural_cycles_figure,
)


def test_halving_phases_follow_specified_progress_bands():
    assert halving_phase(0.0) == "EARLY_EXPANSION"
    assert halving_phase(20.0) == "LATE_EXPANSION"
    assert halving_phase(35.0) == "CYCLE_TOP_RISK"
    assert halving_phase(42.0) == "BEAR_DELEVERAGING"
    assert halving_phase(55.0) == "BOTTOMING_TRANSITION"
    assert halving_phase(70.0) == "ACCUMULATION_PRE_HALVING"


def test_halving_calendar_uses_actual_events_and_projected_april_2028():
    prior = halving_cycle_position("2024-04-19")
    current = halving_cycle_position("2026-09-25")

    assert prior["previous_halving"] == pd.Timestamp("2020-05-11")
    assert prior["next_halving"] == pd.Timestamp("2024-04-20")
    assert current["previous_halving"] == pd.Timestamp("2024-04-20")
    assert current["next_halving"] == BTC_PROJECTED_HALVING
    assert current["next_halving_projected"] is True
    assert current["phase"] == "BOTTOMING_TRANSITION"


def test_liquidity_phase_uses_only_troughs_known_at_signal_date():
    before_2019_anchor = liquidity_cycle_position("2018-10-01", "2018-10-01")
    after_2019_anchor = liquidity_cycle_position("2019-04-01", "2019-04-01")
    forecast_before_2019_anchor = liquidity_cycle_position("2020-01-01", "2018-10-01")

    assert before_2019_anchor["anchor"] == pd.Timestamp("2015-03-06")
    assert after_2019_anchor["anchor"] == pd.Timestamp("2019-03-08")
    assert forecast_before_2019_anchor["anchor"] == pd.Timestamp("2019-07-27")
    assert forecast_before_2019_anchor["anchor"] != pd.Timestamp("2019-03-08")


def test_btc_next_cycle_range_runs_until_next_accumulation_pre_halving_phase():
    btc, canonical_cycle, macro = _cycle_inputs()
    history = build_btc_cycle_history(btc, canonical_cycle, macro)
    latest = history.loc[~history["Projected"], "Date"].max()
    start, end, include_forecast = btc_cycle_time_range(history, "Next Cycle")

    assert BTC_RANGE_OPTIONS == ("1Y", "3Y", "5Y", "10Y", "MAX", "Next Cycle")
    assert start == latest - pd.DateOffset(months=12)
    assert end == next_accumulation_pre_halving_start(latest)
    assert halving_cycle_position(end)["phase"] == "ACCUMULATION_PRE_HALVING"
    assert include_forecast is True


def test_btc_gold_ratio_uses_latest_gold_close_without_lookahead():
    btc = pd.DataFrame(
        {
            "date": pd.to_datetime(["2024-01-12", "2024-01-19", "2024-01-26"]),
            "Close": [42000.0, 43000.0, 44000.0],
        }
    )
    gold = pd.Series(
        [2000.0, 2050.0, 2100.0],
        index=pd.to_datetime(["2024-01-10", "2024-01-17", "2024-01-24"]),
        name="gold_price",
    )

    ratio = build_btc_gold_ratio_history(btc, gold)

    assert ratio["Gold_Price"].tolist() == [2000.0, 2050.0, 2100.0]
    np.testing.assert_allclose(ratio["BTC_GOLD_Ratio"], [21.0, 43000.0 / 2050.0, 44000.0 / 2100.0])


def test_btc_mcp_history_extends_to_index_btcusd_start_without_overlapping_yahoo():
    tradingview = pd.DataFrame(
        {
            "date": pd.DatetimeIndex(["2009-10-05", "2014-09-08", "2014-09-15"], dtype="datetime64[ns]"),
            "open": [0.00076, 475.0, 399.0],
            "high": [0.00120, 480.0, 410.0],
            "low": [0.00076, 470.0, 375.0],
            "close": [0.00115, 475.49, 375.35],
            "volume": [0.0, 1000.0, 1200.0],
        }
    )
    yahoo_dates = pd.DatetimeIndex(
        ["2014-09-17", "2014-09-18", "2014-09-19", "2014-09-20"], dtype="datetime64[ns]"
    )
    yahoo = pd.DataFrame(
        {
            "Open": [457.0, 457.0, 424.0, 394.8],
            "High": [468.0, 457.0, 428.0, 410.0],
            "Low": [452.0, 413.0, 384.0, 390.0],
            "Close": [457.33, 424.44, 394.80, 404.42],
            "Volume": [100.0, 200.0, 300.0, 400.0],
        },
        index=yahoo_dates,
    )

    history = merge_btc_mcp_weekly_history(yahoo, tradingview, today="2014-09-21")

    assert history.iloc[0]["date"] == pd.Timestamp("2009-10-11")
    assert history.iloc[0]["Close"] == 0.00115
    assert pd.Timestamp("2014-09-14") in set(history["date"])
    assert pd.Timestamp("2014-09-21") not in set(history["date"])
    assert pd.Timestamp("2014-09-26") not in set(history["date"])
    assert history.loc[history["date"].eq("2014-09-19"), "Close"].iloc[0] == 394.80


def test_btc_gold_ratio_chart_obeys_shared_btc_time_range():
    btc, canonical_cycle, macro = _cycle_inputs()
    history = build_btc_cycle_history(btc, canonical_cycle, macro)
    gold_dates = pd.to_datetime(btc["date"])
    gold = pd.Series(np.linspace(1200.0, 2800.0, len(gold_dates)), index=gold_dates, name="gold_price")
    ratio = build_btc_gold_ratio_history(btc, gold)
    latest_observation = history.loc[~history["Projected"], "Date"].max()

    for choice in ("1Y", "5Y", "Next Cycle"):
        start, end, _ = btc_cycle_time_range(history, choice)
        figure = _build_btc_gold_ratio_figure(history, ratio, choice)

        assert figure.layout.xaxis.range == (start, end)
        assert pd.to_datetime(figure.data[0].x).min() >= start
        assert pd.to_datetime(figure.data[0].x).max() <= latest_observation
        assert figure.data[0].name == "BTC / Gold"


def test_btc_etf_flow_chart_obeys_shared_btc_time_range():
    btc, canonical_cycle, macro = _cycle_inputs()
    history = build_btc_cycle_history(btc, canonical_cycle, macro)
    flow_history = pd.DataFrame(
        {
            "date": btc["date"],
            "ETF_Flow_Intensity_4W": np.sin(np.arange(len(btc)) / 9.0),
            "ETF_Flow_3Y_Pctl": np.linspace(0.0, 100.0, len(btc)),
        }
    )
    latest_observation = history.loc[~history["Projected"], "Date"].max()

    for choice in ("1Y", "5Y", "MAX", "Next Cycle"):
        start, end, _ = btc_cycle_time_range(history, choice)
        figure = _build_btc_etf_flow_intensity_figure(history, flow_history, choice)

        assert figure.layout.xaxis.range == (start, end)
        assert [trace.name for trace in figure.data] == ["4W Flow Intensity", "3Y Percentile"]
        assert figure.data[1].yaxis == "y2"
        assert figure.layout.yaxis2.range == (0, 100)
        assert pd.to_datetime(figure.data[0].x).min() >= start
        assert pd.to_datetime(figure.data[0].x).max() <= latest_observation


def test_secondary_percentile_is_trailing_and_point_in_time():
    dates = pd.date_range("2020-01-03", periods=180, freq="W-FRI")
    values = pd.Series(np.sin(np.arange(len(dates)) / 8.0), index=dates)
    original = point_in_time_percentile(values)
    changed = values.copy()
    changed.iloc[-1] = 1000.0
    changed_rank = point_in_time_percentile(changed)

    assert np.isnan(original.iloc[102])
    assert np.isfinite(original.iloc[103])
    assert np.isclose(original.iloc[-2], changed_rank.iloc[-2])
    assert changed_rank.iloc[-1] > original.iloc[-1]


def _cycle_inputs():
    btc_dates = pd.date_range("2014-01-03", "2026-09-25", freq="W-FRI")
    index = np.arange(len(btc_dates), dtype=float)
    btc = pd.DataFrame({"date": btc_dates, "Close": 1000.0 * np.exp(index / 230.0)})
    macro = pd.DataFrame(
        {
            "Date": btc_dates,
            "BTC_DXY_Close": 100.0 + np.sin(index / 12.0) * 4 + index / 500.0,
            "BTC_US2Y": 3.0 + np.sin(index / 17.0),
            "BTC_RealYield": 1.0 + np.cos(index / 23.0) * 0.6,
        }
    )
    m2_dates = pd.date_range("1990-01-01", "2026-08-01", freq="MS")
    m2_index = np.arange(len(m2_dates), dtype=float)
    cycle = pd.DataFrame({"Date": m2_dates, "PrimaryMarketCycle": np.sin(m2_index / 16.0)})
    return btc, cycle, macro


def test_btc_cycle_reconciles_scores_and_has_phase_only_projection():
    btc, canonical_cycle, macro = _cycle_inputs()
    history = build_btc_cycle_history(btc, canonical_cycle, macro)
    current = history.loc[~history["Projected"]].iloc[-1]
    projected = history.loc[history["Projected"]]

    assert not btc_cycle_validation(history)
    assert current["Liquidity_Trough_Anchor"] == pd.Timestamp("2022-10-28")
    available_m2 = canonical_cycle.loc[canonical_cycle["Date"].le(current["Date"])].iloc[-1]
    assert np.isclose(current["GlobalM2PrimaryCycle"], available_m2["PrimaryMarketCycle"])
    for horizon in ("3M", "6M", "9M", "12M"):
        inputs = (
            current[f"HalvingBase_{horizon}"]
            + current[f"LiquidityCycleModifier_{horizon}"]
            + current[f"SecondaryMacroModifier_{horizon}"]
        )
        assert np.isclose(current[f"BTC_MACRO_{horizon}"], np.clip(inputs, 0.0, 95.0))
        assert projected[f"SecondaryMacroModifier_{horizon}"].eq(0.0).all()
    assert projected["BTC_Price"].isna().all()
    assert projected["Historical_Projected_Flag"].eq("PROJECTED").all()
    assert projected["Date"].max() <= next_accumulation_pre_halving_start(current["Date"])
    assert BTC_LIQUIDITY_CYCLE_MONTHS == 52.7
    assert projected["GlobalM2CycleProjected"].any()
    assert projected.loc[projected["GlobalM2CycleProjected"], "Date"].min() > canonical_cycle["Date"].max()
    projected_cycle = projected.loc[projected["GlobalM2CycleProjected"]]
    next_trough = liquidity_cycle_position(current["Date"], current["Date"])["next_trough"]
    closest_trough = projected_cycle.loc[(projected_cycle["Date"] - next_trough).abs().idxmin()]
    assert abs((closest_trough["Date"] - next_trough).days) <= 4
    trough_window = projected_cycle.loc[
        projected_cycle["Date"].between(
            next_trough - pd.Timedelta(45, unit="D"),
            next_trough + pd.Timedelta(45, unit="D"),
        )
    ]
    assert closest_trough["GlobalM2PrimaryCycle"] <= trough_window["GlobalM2PrimaryCycle"].min() + 0.02


def test_btc_cycle_xlsx_contains_full_history_and_projected_rows():
    btc, canonical_cycle, macro = _cycle_inputs()
    history = build_btc_cycle_history(btc, canonical_cycle, macro)

    workbook = load_workbook(BytesIO(btc_cycle_export_xlsx(history)), read_only=True)

    assert workbook.sheetnames == ["BTC Cycle"]
    sheet = workbook["BTC Cycle"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert "BTC_MACRO_3M" in headers
    assert "Historical_Projected_Flag" in headers
    assert sheet.max_row == len(history) + 1


def test_modular_cycle_forecast_uses_two_historical_modules_and_canonical_start_price():
    btc, _, _ = _cycle_inputs()
    target = 262_469.205

    forecast = build_btc_modular_cycle_forecast(btc, target, "Base")
    source_dates = pd.to_datetime(btc["date"])
    nearest_start = (source_dates - BTC_MODULAR_CYCLE_START).abs().idxmin()
    expected_duration = np.mean(
        [
            (
                source_dates.iloc[(source_dates - end).abs().argmin()]
                - source_dates.iloc[(source_dates - start).abs().argmin()]
            ).days
            / 7
            for start, end in BTC_MODULAR_HISTORICAL_MODULES
        ]
    )

    assert forecast["Date"].iloc[0] == BTC_MODULAR_CYCLE_START
    assert np.isclose(forecast["NextCycle_StartPrice"].iloc[0], btc.loc[nearest_start, "Close"])
    assert np.isclose(forecast["NextCycle_DurationWeeks"].iloc[0], expected_duration)
    assert np.isclose(expected_duration, 203.5)
    assert forecast["Date"].iloc[-1] == forecast["NextCycle_EndDate"].iloc[0]
    assert np.isclose(forecast["NextCycle_ProgressPct"].iloc[0], 0.0)
    assert np.isclose(forecast["NextCycle_ProgressPct"].iloc[-1], 100.0)
    assert np.isclose(forecast["NextCycle_ModelPrice"].max(), target, rtol=0, atol=1e-7)
    assert forecast["NextCycle_PeakDate"].nunique() == 1
    assert forecast["NextCycle_HistoricalModule1_Start"].iloc[0] == pd.Timestamp("2018-11-02")
    assert forecast["NextCycle_HistoricalModule1_End"].iloc[0] == pd.Timestamp("2022-11-25")
    assert forecast["NextCycle_HistoricalModule2_Start"].iloc[0] == pd.Timestamp("2022-09-27")
    assert forecast["NextCycle_HistoricalModule2_End"].iloc[0] == pd.Timestamp("2026-06-26")
    assert forecast["ProjectedFlag"].all()


def test_modular_cycle_target_change_rescales_amplitude_without_changing_timing():
    btc, _, _ = _cycle_inputs()
    base = build_btc_modular_cycle_forecast(btc, 262_469.205, "Base")
    strong = build_btc_modular_cycle_forecast(btc, 310_000, "Strong Liquidity")

    assert base["Date"].equals(strong["Date"])
    assert np.allclose(base["NextCycle_ProgressPct"], strong["NextCycle_ProgressPct"])
    assert base["NextCycle_PeakDate"].iloc[0] == strong["NextCycle_PeakDate"].iloc[0]
    base_shape = np.log(base["NextCycle_ModelPrice"] / base["NextCycle_StartPrice"]) / np.log(
        base["NextCycle_TargetPeak"] / base["NextCycle_StartPrice"]
    )
    strong_shape = np.log(strong["NextCycle_ModelPrice"] / strong["NextCycle_StartPrice"]) / np.log(
        strong["NextCycle_TargetPeak"] / strong["NextCycle_StartPrice"]
    )
    assert np.allclose(base_shape, strong_shape)


def test_next_cycle_chart_overlays_dashed_model_and_extends_only_price_chart():
    btc, canonical_cycle, macro = _cycle_inputs()
    history = build_btc_cycle_history(btc, canonical_cycle, macro)
    target = 262_469.205
    forecast = build_btc_modular_cycle_forecast(btc, target, "Base")
    latest = history.loc[~history["Projected"], "Date"].max()

    next_cycle_fig = _build_btc_price_halving_figure(history, "Next Cycle", forecast)
    standard_fig = _build_btc_price_halving_figure(history, "5Y", forecast)

    assert next_cycle_fig.layout.xaxis.range[1] == forecast["NextCycle_EndDate"].iloc[0]
    assert pd.to_datetime(next_cycle_fig.data[0].x).max() == latest
    model_trace = next(trace for trace in next_cycle_fig.data if trace.name == "Next Cycle Model — Base")
    assert model_trace.line.dash == "dash"
    assert pd.to_datetime(model_trace.x).min() == BTC_MODULAR_CYCLE_START
    assert pd.to_datetime(model_trace.x).max() == forecast["NextCycle_EndDate"].iloc[0]
    assert np.isclose(max(model_trace.y), target, rtol=0, atol=1e-7)
    assert standard_fig.layout.xaxis.range[1] == latest
    assert not any("Next Cycle Model" in str(trace.name) for trace in standard_fig.data)


def test_btc_cycle_export_adds_weekly_next_cycle_forecast_sheet():
    btc, canonical_cycle, macro = _cycle_inputs()
    history = build_btc_cycle_history(btc, canonical_cycle, macro)
    forecast = build_btc_modular_cycle_forecast(btc, 262_469.205, "Base")
    workbook = load_workbook(BytesIO(btc_cycle_export_xlsx(history, forecast)), read_only=True)

    assert workbook.sheetnames == ["BTC Cycle", "Next Cycle Forecast"]
    sheet = workbook["Next Cycle Forecast"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert {
        "Date",
        "NextCycle_ModelPrice",
        "NextCycle_ModelMultiple",
        "NextCycle_ProgressPct",
        "NextCycle_TargetPeak",
        "NextCycle_Scenario",
        "NextCycle_PeakDate",
        "NextCycle_EndDate",
        "HalvingPhase",
        "ProjectedFlag",
    }.issubset(headers)
    assert sheet.max_row == len(forecast) + 1


def test_btc_cycle_charts_separate_observed_price_from_projected_scores():
    btc, canonical_cycle, macro = _cycle_inputs()
    liquidity_score = pd.DataFrame(
        {
            "date": btc["date"],
            "global_liquidity_score": np.linspace(20.0, 80.0, len(btc)),
        }
    )
    history = build_btc_cycle_history(btc, canonical_cycle, macro, global_liquidity_score=liquidity_score)
    latest_observation = history.loc[~history["Projected"], "Date"].max()

    price_fig = _build_btc_price_halving_figure(history, "Next Cycle")
    liquidity_fig = _build_structural_cycles_figure(history, "Next Cycle")
    score_fig = _build_macro_score_figure(history, "3M", "Next Cycle")

    assert pd.to_datetime(price_fig.data[0].x).max() == latest_observation
    assert price_fig.layout.xaxis.range[1] == next_accumulation_pre_halving_start(latest_observation)
    assert price_fig.data[1].name == "Global Liquidity Score"
    assert price_fig.data[1].yaxis == "y2"
    assert price_fig.layout.yaxis2.side == "right"
    assert [trace.name for trace in liquidity_fig.data] == [
        "Global M2 Primary Liquidity Cycle",
        "Projected Global M2 Cycle",
    ]
    assert len(score_fig.data) == 2
    assert score_fig.data[1].line.dash == "dash"
    assert pd.to_datetime(score_fig.data[1].x).max() <= next_accumulation_pre_halving_start(latest_observation)


def test_standard_btc_ranges_exclude_forecasts_and_end_at_latest_observation():
    btc, canonical_cycle, macro = _cycle_inputs()
    liquidity_score = pd.DataFrame(
        {
            "date": btc["date"],
            "global_liquidity_score": np.linspace(20.0, 80.0, len(btc)),
        }
    )
    history = build_btc_cycle_history(btc, canonical_cycle, macro, global_liquidity_score=liquidity_score)
    latest = history.loc[~history["Projected"], "Date"].max()

    price_fig = _build_btc_price_halving_figure(history, "5Y")
    cycle_fig = _build_structural_cycles_figure(history, "5Y")
    score_fig = _build_macro_score_figure(history, "3M", "5Y")

    assert price_fig.layout.xaxis.range[1] == latest
    assert any(trace.name == "Global Liquidity Score" for trace in price_fig.data)
    assert [trace.name for trace in cycle_fig.data] == ["Global M2 Primary Liquidity Cycle"]
    assert len(score_fig.data) == 1
    assert pd.to_datetime(score_fig.data[0].x).max() == latest
