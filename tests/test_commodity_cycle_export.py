from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from commodity_cycle.export import commodity_tables_to_xlsx


def test_commodity_tables_export_to_two_formatted_xlsx_sheets():
    primary = pd.DataFrame({
        "Sector": ["Energy", "Energy"],
        "Commodity": ["WTI", "Natural Gas"],
        "Price": [90.421, 3.026],
        "Return 3M": [0.30123, -0.07601],
        "Raw Curve State": ["Backwardation", "Contango"],
        "COT 5Y Percentile": [20.3846, 10.0],
        "Spread 5Y Seasonal Percentile": [20.0, 80.0],
        "Price Date": pd.to_datetime(["2026-09-30", "2026-09-30"]),
    })
    diagnostics = pd.DataFrame({
        "Sector": ["Energy"], "Commodity": ["WTI"], "COT 3Y Percentile": [10.1234],
        "Spread 10Y Seasonal Percentile": [30.0],
    })

    workbook_bytes = commodity_tables_to_xlsx(primary, diagnostics)
    workbook = load_workbook(BytesIO(workbook_bytes))

    assert workbook.sheetnames == ["Confirmation", "Diagnostics"]
    confirmation = workbook["Confirmation"]
    diagnostics_sheet = workbook["Diagnostics"]
    assert [cell.value for cell in confirmation[1]] == primary.columns.tolist()
    assert [cell.value for cell in diagnostics_sheet[1]] == diagnostics.columns.tolist()
    assert confirmation["C2"].number_format == "0.00"
    assert confirmation["D2"].number_format == "0.00%"
    assert confirmation["H2"].number_format == "yyyy-mm-dd"
    formatting = list(confirmation.conditional_formatting)
    assert len(formatting) == 2
    assert sorted(len(item.rules) for item in formatting) == [1, 2]
